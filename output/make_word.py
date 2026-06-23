# -*- coding: utf-8 -*-
"""Создание документа Word с заданиями 4-5.

Содержит:
  * формулы (OMML) - произведение матриц и определённый интеграл;
  * таблицу "Площадь континентов" по образцу;
  * НАТИВНУЮ (редактируемую) круговую диаграмму Office, а не картинку;
  * метаданные: Авторы и "Кем сохранён".
"""
import io
import os
import re
import shutil
import zipfile

from openpyxl import Workbook as XlWorkbook

from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

M = "http://schemas.openxmlformats.org/officeDocument/2006/math"
AUTHOR = "Студент 9/25. РПО"
MARKER = "CHART_PLACEHOLDER_XYZ"

continents = [
    ("Европа", 10.00),
    ("Азия", 43.40),
    ("Северная Америка", 24.71),
    ("Южная Америка", 17.84),
    ("Австралия", 7.66),
    ("Антарктида", 14.10),
    ("Африка", 30.30),
]
TOTAL = 148.01


# ----------------------------------------------------------------- OMML helpers
def omath_wrap(inner):
    return ('<m:oMathPara xmlns:m="%s"><m:oMath>%s</m:oMath></m:oMathPara>'
            % (M, inner))


def run(text):
    return "<m:r><m:t>%s</m:t></m:r>" % text


def matrix(rows):
    mr = ""
    for row in rows:
        cells = "".join("<m:e>%s</m:e>" % run(v) for v in row)
        mr += "<m:mr>%s</m:mr>" % cells
    return "<m:m>%s</m:m>" % mr


def bracket(inner):
    return ('<m:d><m:dPr><m:begChr m:val="("/><m:endChr m:val=")"/></m:dPr>'
            '<m:e>%s</m:e></m:d>') % inner


def add_math(doc, inner):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p._p.append(parse_xml(omath_wrap(inner)))
    return p


# ----------------------------------------------------------------- build doc
doc = Document()
doc.styles["Normal"].font.name = "Times New Roman"
doc.styles["Normal"].font.size = Pt(14)

cp = doc.core_properties
cp.author = AUTHOR            # Авторы
cp.last_modified_by = AUTHOR  # Кем сохранён
cp.title = "Задания в Word"
cp.category = "РПО"

h = doc.add_paragraph()
hr = h.add_run("Выполнить в Word")
hr.bold = True
hr.font.size = Pt(14)

# -------- Задание 4
p = doc.add_paragraph()
p.add_run("Задание 4").italic = True
doc.add_paragraph("Написать формулы:")

mat_a = matrix([["1", "2", "4"], ["3", "7", "8"], ["5", "6", "9"]])
mat_b = matrix([["1"], ["5"], ["9"]])
add_math(doc, bracket(mat_a) + run("∙") + bracket(mat_b))

doc.add_paragraph("")

x_sq = "<m:sSup><m:e>%s</m:e><m:sup>%s</m:sup></m:sSup>" % (run("x"), run("2"))
num = "<m:num>%s</m:num>" % run("cos(0.8∙x+1.2)")
den = "<m:den>%s%s%s</m:den>" % (run("1.5∙x+sin("), x_sq, run("+0.6)"))
frac = "<m:f>%s%s</m:f>" % (num, den)
nary = ('<m:nary><m:naryPr><m:chr m:val="∫"/><m:limLoc m:val="subSup"/>'
        '</m:naryPr><m:sub>%s</m:sub><m:sup>%s</m:sup><m:e>%s%s</m:e></m:nary>'
        % (run("0.8"), run("0.9"), frac, run("dx")))
add_math(doc, nary)

doc.add_paragraph("")

# -------- Задание 5
p = doc.add_paragraph()
p.add_run("Задание 5").italic = True
doc.add_paragraph("Создайте и оформите таблицу по образцу. "
                  "Построить круговую диаграмму по полученным данным")

table = doc.add_table(rows=0, cols=2)
table.style = "Table Grid"
table.alignment = WD_TABLE_ALIGNMENT.CENTER


def set_cell(cell, text, bold=False, center=True):
    cell.text = ""
    para = cell.paragraphs[0]
    if center:
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = para.add_run(text)
    r.bold = bold
    r.font.name = "Times New Roman"
    r.font.size = Pt(14)


def shade(cell, color="D9D9D9"):
    cell._tc.get_or_add_tcPr().append(
        parse_xml(r'<w:shd %s w:fill="%s"/>' % (nsdecls("w"), color)))


hdr = table.add_row().cells
a = hdr[0].merge(hdr[1])
set_cell(a, "Площадь континентов", bold=True)
shade(a)

sub = table.add_row().cells
set_cell(sub[0], "Название континента", bold=True)
set_cell(sub[1], "Площадь (млн.кв.км)", bold=True)
shade(sub[0]); shade(sub[1])

for name, area in continents:
    cells = table.add_row().cells
    set_cell(cells[0], name, center=False)
    set_cell(cells[1], "%.2f" % area)

cells = table.add_row().cells
set_cell(cells[0], "Всего:", bold=True, center=False)
set_cell(cells[1], "%.2f" % TOTAL, bold=True)

doc.add_paragraph("")

# placeholder paragraph для нативной диаграммы
ph = doc.add_paragraph()
ph.alignment = WD_ALIGN_PARAGRAPH.CENTER
ph.add_run(MARKER)

tmp_path = "output/_tmp_word.docx"
doc.save(tmp_path)


# ----------------------------------------------------------------- native chart
def build_embedded_xlsx():
    wb = XlWorkbook()
    ws = wb.active
    ws.title = "Лист1"
    ws["A1"] = "Континент"
    ws["B1"] = "Площадь (млн.кв.км)"
    for i, (name, area) in enumerate(continents, start=2):
        ws["A%d" % i] = name
        ws["B%d" % i] = area
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def str_cache(values):
    pts = "".join('<c:pt idx="%d"><c:v>%s</c:v></c:pt>' % (i, v)
                  for i, v in enumerate(values))
    return ('<c:strCache><c:ptCount val="%d"/>%s</c:strCache>'
            % (len(values), pts))


def num_cache(values):
    pts = "".join('<c:pt idx="%d"><c:v>%s</c:v></c:pt>' % (i, v)
                  for i, v in enumerate(values))
    return ('<c:numCache><c:formatCode>General</c:formatCode>'
            '<c:ptCount val="%d"/>%s</c:numCache>' % (len(values), pts))


def build_chart_xml():
    names = [c[0] for c in continents]
    vals = [c[1] for c in continents]
    n = len(continents)
    cat_ref = ("<c:cat><c:strRef><c:f>Лист1!$A$2:$A$%d</c:f>%s</c:strRef></c:cat>"
               % (n + 1, str_cache(names)))
    val_ref = ("<c:val><c:numRef><c:f>Лист1!$B$2:$B$%d</c:f>%s</c:numRef></c:val>"
               % (n + 1, num_cache(vals)))
    ser = (
        '<c:ser><c:idx val="0"/><c:order val="0"/>'
        '<c:tx><c:strRef><c:f>Лист1!$B$1</c:f>'
        '<c:strCache><c:ptCount val="1"/><c:pt idx="0">'
        '<c:v>Площадь (млн.кв.км)</c:v></c:pt></c:strCache></c:strRef></c:tx>'
        '<c:dLbls><c:showLegendKey val="0"/><c:showVal val="0"/>'
        '<c:showCatName val="0"/><c:showSerName val="0"/>'
        '<c:showPercent val="1"/><c:showBubbleSize val="0"/></c:dLbls>'
        + cat_ref + val_ref + '</c:ser>'
    )
    return (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<c:chartSpace '
        'xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart" '
        'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
        'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
        '<c:chart>'
        '<c:title><c:tx><c:rich><a:bodyPr/><a:p><a:r>'
        '<a:t>Площадь континентов (млн.кв.км)</a:t></a:r></a:p></c:rich></c:tx>'
        '<c:overlay val="0"/></c:title>'
        '<c:autoTitleDeleted val="0"/>'
        '<c:plotArea><c:layout/>'
        '<c:pieChart><c:varyColors val="1"/>'
        + ser +
        '<c:firstSliceAng val="0"/></c:pieChart></c:plotArea>'
        '<c:legend><c:legendPos val="r"/><c:overlay val="0"/></c:legend>'
        '<c:plotVisOnly val="1"/><c:dispBlanksAs val="gap"/>'
        '</c:chart>'
        '<c:externalData r:id="rId1"><c:autoUpdate val="0"/></c:externalData>'
        '</c:chartSpace>'
    )


CHART_RID = "rId900"
drawing = (
    '<w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
    '<w:drawing '
    'xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
    'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
    'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
    'xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart">'
    '<wp:inline distT="0" distB="0" distL="0" distR="0">'
    '<wp:extent cx="5486400" cy="3429000"/>'
    '<wp:effectExtent l="0" t="0" r="0" b="0"/>'
    '<wp:docPr id="100" name="Диаграмма 1"/>'
    '<wp:cNvGraphicFramePr/>'
    '<a:graphic><a:graphicData '
    'uri="http://schemas.openxmlformats.org/drawingml/2006/chart">'
    '<c:chart r:id="%s"/></a:graphicData></a:graphic>'
    '</wp:inline></w:drawing></w:r>' % CHART_RID
)


def inject_native_chart(src, dst):
    chart_xml = build_chart_xml()
    xlsx_bytes = build_embedded_xlsx()
    chart_rels = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rId1" '
        'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/package" '
        'Target="../embeddings/Microsoft_Excel_Worksheet1.xlsx"/>'
        '</Relationships>'
    )

    zin = zipfile.ZipFile(src, "r")
    names = zin.namelist()

    # document.xml: заменяем маркер на нативную диаграмму
    document = zin.read("word/document.xml").decode("utf-8")
    marker_run = re.search(r"<w:r>(?:(?!</w:r>).)*?%s.*?</w:r>" % MARKER,
                           document, re.S)
    assert marker_run, "marker run not found"
    document = document.replace(marker_run.group(0), drawing)

    # document.xml.rels: добавляем связь с диаграммой
    rels = zin.read("word/_rels/document.xml.rels").decode("utf-8")
    new_rel = ('<Relationship Id="%s" '
               'Type="http://schemas.openxmlformats.org/officeDocument/2006/'
               'relationships/chart" Target="charts/chart1.xml"/>' % CHART_RID)
    rels = rels.replace("</Relationships>", new_rel + "</Relationships>")

    # [Content_Types].xml
    ct = zin.read("[Content_Types].xml").decode("utf-8")
    additions = ""
    if 'Extension="xlsx"' not in ct:
        additions += ('<Default Extension="xlsx" ContentType="application/vnd.'
                      'openxmlformats-officedocument.spreadsheetml.sheet"/>')
    additions += ('<Override PartName="/word/charts/chart1.xml" '
                  'ContentType="application/vnd.openxmlformats-officedocument.'
                  'drawingml.chart+xml"/>')
    ct = ct.replace("</Types>", additions + "</Types>")

    zout = zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED)
    for item in names:
        if item == "word/document.xml":
            zout.writestr(item, document)
        elif item == "word/_rels/document.xml.rels":
            zout.writestr(item, rels)
        elif item == "[Content_Types].xml":
            zout.writestr(item, ct)
        else:
            zout.writestr(item, zin.read(item))
    # новые части
    zout.writestr("word/charts/chart1.xml", chart_xml)
    zout.writestr("word/charts/_rels/chart1.xml.rels", chart_rels)
    zout.writestr("word/embeddings/Microsoft_Excel_Worksheet1.xlsx", xlsx_bytes)
    zout.close()
    zin.close()


inject_native_chart(tmp_path, "output/Задания_Word.docx")
os.remove(tmp_path)
print("Word saved (native chart)")
